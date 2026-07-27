import os
import csv
from datetime import datetime

# PDF generation using ReportLab
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# Excel generation using openpyxl
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "reports")

def ensure_reports_dir():
    os.makedirs(REPORTS_DIR, exist_ok=True)
    return REPORTS_DIR

def generate_pdf_report(job_title, job_category, results_list, filename_prefix="Candidate_Report"):
    """
    Generates a professional PDF candidate evaluation report using ReportLab.
    """
    ensure_reports_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    pdf_filename = f"{filename_prefix}_{timestamp}.pdf"
    pdf_path = os.path.join(REPORTS_DIR, pdf_filename)
    
    doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1e1b4b"),
        fontName="Helvetica-Bold",
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'SubTitle',
        parent=styles['Normal'],
        fontSize=11,
        textColor=colors.HexColor("#475569"),
        fontName="Helvetica",
        spaceAfter=15
    )
    
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#1e293b")
    )
    
    bold_cell = ParagraphStyle(
        'BoldCellText',
        parent=cell_style,
        fontName="Helvetica-Bold"
    )

    elements = []
    
    elements.append(Paragraph("AI Resume Screening - Candidate Ranking Report", title_style))
    elements.append(Paragraph(f"<b>Job Title:</b> {job_title} | <b>Category:</b> {job_category} | <b>Generated:</b> {datetime.now().strftime('%B %d, %Y %I:%M %p')}", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#4f46e5"), spaceAfter=15))
    
    # Table headers
    headers = ["Rank", "Candidate Name", "Email", "Matching Score", "Cosine Sim", "Skill Match", "Status"]
    table_data = [[Paragraph(h, bold_cell) for h in headers]]
    
    for rank, item in enumerate(results_list, 1):
        name = item.get("parsed_name", "N/A")
        email = item.get("parsed_email", "N/A")
        score = f"{item.get('matching_score', 0):.1f}%"
        cosine = f"{item.get('cosine_similarity', 0)*100:.1f}%"
        skill = f"{item.get('skill_match_score', 0):.1f}%"
        status = item.get("status", "Evaluated")
        
        status_color = "#15803d" if "Highly" in status else ("#b45309" if "Recommended" in status else "#b91c1c")
        status_p = Paragraph(f"<font color='{status_color}'><b>{status}</b></font>", cell_style)
        
        row = [
            Paragraph(str(rank), cell_style),
            Paragraph(name, bold_cell),
            Paragraph(email, cell_style),
            Paragraph(score, bold_cell),
            Paragraph(cosine, cell_style),
            Paragraph(skill, cell_style),
            status_p
        ]
        table_data.append(row)
        
    t = Table(table_data, colWidths=[35, 110, 140, 75, 65, 65, 90])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")])
    ]))
    
    elements.append(t)
    elements.append(Spacer(1, 20))
    elements.append(Paragraph("<i>Report generated automatically by AI Resume Screening Engine. Confidential.</i>", subtitle_style))
    
    doc.build(elements)
    return pdf_filename, pdf_path

def generate_excel_report(job_title, job_category, results_list, filename_prefix="Candidate_Report"):
    """
    Generates a styled Excel spreadsheet report using OpenPyXL.
    """
    ensure_reports_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_filename = f"{filename_prefix}_{timestamp}.xlsx"
    excel_path = os.path.join(REPORTS_DIR, excel_filename)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candidate Ranking"
    
    # Styles
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E1B4B")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # Header Information
    ws.cell(row=1, column=1, value="AI Resume Screening - Candidate Ranking").font = title_font
    ws.cell(row=2, column=1, value=f"Job Title: {job_title} | Category: {job_category} | Exported: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = regular_font
    
    headers = ["Rank", "Candidate Name", "Email", "Phone", "Overall Match Score (%)", "Cosine Similarity (%)", "Skill Match (%)", "ML Confidence (%)", "Missing Skills", "Status"]
    
    start_row = 4
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for rank, item in enumerate(results_list, 1):
        r = start_row + rank
        ws.cell(row=r, column=1, value=rank).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=item.get("parsed_name", "N/A")).font = bold_font
        ws.cell(row=r, column=3, value=item.get("parsed_email", "N/A"))
        ws.cell(row=r, column=4, value=item.get("parsed_phone", "N/A"))
        ws.cell(row=r, column=5, value=round(item.get("matching_score", 0), 1)).font = bold_font
        ws.cell(row=r, column=6, value=round(item.get("cosine_similarity", 0) * 100, 1))
        ws.cell(row=r, column=7, value=round(item.get("skill_match_score", 0), 1))
        ws.cell(row=r, column=8, value=round(item.get("ml_probability", 0) * 100, 1))
        ws.cell(row=r, column=9, value=item.get("missing_skills", "None"))
        
        status_cell = ws.cell(row=r, column=10, value=item.get("status", "N/A"))
        status = item.get("status", "")
        if "Highly" in status:
            status_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            status_cell.font = Font(name="Calibri", size=11, bold=True, color="15803D")
        elif "Recommended" in status:
            status_cell.fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
            status_cell.font = Font(name="Calibri", size=11, bold=True, color="A16207")
        else:
            status_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            status_cell.font = Font(name="Calibri", size=11, bold=True, color="B91C1C")
            
        for col_num in range(1, len(headers) + 1):
            ws.cell(row=r, column=col_num).border = thin_border
            
    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    wb.save(excel_path)
    return excel_filename, excel_path

def generate_csv_report(job_title, results_list, filename_prefix="Candidate_Report"):
    """
    Generates a CSV export of candidates.
    """
    ensure_reports_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    csv_filename = f"{filename_prefix}_{timestamp}.csv"
    csv_path = os.path.join(REPORTS_DIR, csv_filename)
    
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Rank", "Name", "Email", "Phone", "Matching Score (%)", "Cosine Sim (%)", "Skill Score (%)", "ML Prob (%)", "Missing Skills", "Status"])
        for rank, item in enumerate(results_list, 1):
            writer.writerow([
                rank,
                item.get("parsed_name", "N/A"),
                item.get("parsed_email", "N/A"),
                item.get("parsed_phone", "N/A"),
                round(item.get("matching_score", 0), 1),
                round(item.get("cosine_similarity", 0) * 100, 1),
                round(item.get("skill_match_score", 0), 1),
                round(item.get("ml_probability", 0) * 100, 1),
                item.get("missing_skills", ""),
                item.get("status", "")
            ])
            
    return csv_filename, csv_path
