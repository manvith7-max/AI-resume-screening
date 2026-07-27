import os
import re

# Comprehensive list of skills for parsing
SKILL_DICTIONARY = [
    # Programming Languages
    "python", "javascript", "typescript", "java", "c++", "c#", "c", "ruby", "php", "swift", "kotlin", "go", "golang",
    "rust", "r", "matlab", "scala", "perl", "bash", "powershell", "sql", "html", "css", "sass", "less",
    
    # Web Frameworks & Libraries
    "react", "react.js", "angular", "vue", "vue.js", "next.js", "nuxt.js", "node.js", "express", "django", "flask",
    "fastapi", "spring", "spring boot", "laravel", "symfony", "asp.net", "bootstrap", "tailwind", "jquery",
    
    # Data Science & Machine Learning
    "machine learning", "deep learning", "nlp", "natural language processing", "computer vision", "tensorflow",
    "pytorch", "keras", "scikit-learn", "pandas", "numpy", "scipy", "opencv", "spacy", "nltk", "huggingface",
    "transformers", "xgboost", "lightgbm", "data analysis", "data mining", "predictive modeling", "llm", "rag",
    
    # Databases & Storage
    "postgresql", "mysql", "sqlite", "mongodb", "redis", "elasticsearch", "dynamodb", "cassandra", "oracle",
    "sql server", "firebase", "supabase", "neo4j", "snowflake", "bigquery",
    
    # Cloud & DevOps
    "aws", "azure", "gcp", "google cloud", "docker", "kubernetes", "terraform", "ansible", "jenkins", "gitlab ci",
    "github actions", "circleci", "helm", "prometheus", "grafana", "linux", "unix", "sysadmin", "devops",
    
    # Soft Skills & Business
    "project management", "agile", "scrum", "jira", "leadership", "communication", "problem solving",
    "critical thinking", "financial analysis", "accounting", "risk management", "marketing", "seo",
    "content strategy", "sales", "crm", "salesforce", "customer service", "human resources", "recruitment",
    
    # Design & Tools
    "figma", "adobe xd", "photoshop", "illustrator", "ui design", "ux design", "graphic design", "wireframing",
    "git", "github", "bitbucket", "vscode", "excel", "powerpoint", "word", "tableau", "power bi"
]

def extract_text_from_pdf(filepath):
    """
    Extracts text from a PDF file using PyMuPDF (fitz) or pdfplumber fallback.
    """
    text = ""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(filepath)
        for page in doc:
            text += page.get_text() + "\n"
        doc.close()
        if text.strip():
            return text
    except Exception:
        pass

    try:
        import pdfplumber
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        if text.strip():
            return text
    except Exception:
        pass

    return text

def extract_text_from_docx(filepath):
    """
    Extracts text from a DOCX file using python-docx.
    """
    text = ""
    try:
        import docx
        doc = docx.Document(filepath)
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                row_text = " ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                text += row_text + "\n"
    except Exception as e:
        print(f"Error reading DOCX {filepath}: {e}")
    return text

def extract_text_from_file(filepath):
    """
    Extracts text based on file extension (.pdf, .docx, .txt).
    """
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(filepath)
    elif ext == ".txt":
        try:
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception:
            return ""
    return ""

def parse_email(text):
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    matches = re.findall(email_pattern, text)
    return matches[0] if matches else "N/A"

def parse_phone(text):
    phone_pattern = r'(?:(?:\+|00)\d{1,3}[\s.-]?)?(?:\(?\d{3}\)?[\s.-]?)?\d{3}[\s.-]?\d{4}'
    matches = re.findall(phone_pattern, text)
    for m in matches:
        cleaned = re.sub(r'[^\d+]', '', m)
        if len(cleaned) >= 10:
            return m.strip()
    return "N/A"

def parse_name(text):
    """
    Heuristic name extraction: inspects first few lines for capitalized names.
    """
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    invalid_keywords = ["resume", "curriculum", "vitae", "cv", "page", "contact", "email", "phone", "profile", "summary"]
    
    for line in lines[:5]:
        if len(line) < 35 and not any(kw in line.lower() for kw in invalid_keywords):
            # Check if line looks like a name (words starting with capital letters)
            words = line.split()
            if 2 <= len(words) <= 4 and all(w[0].isupper() for w in words if w.isalpha()):
                return line
    
    # Fallback to first line if reasonably short
    if lines:
        return lines[0][:30]
    return "Candidate"

def parse_skills(text):
    """
    Extracts matched skills from SKILL_DICTIONARY.
    """
    text_lower = text.lower()
    found_skills = set()
    
    for skill in SKILL_DICTIONARY:
        # Use regex boundary checking to avoid partial string matches
        pattern = r'\b' + re.escape(skill) + r'\b'
        if re.search(pattern, text_lower):
            found_skills.add(skill.title() if len(skill) > 3 else skill.upper())
            
    return sorted(list(found_skills))

def parse_education(text):
    """
    Extracts degree keywords and institutions.
    """
    degrees = ["bachelor", "master", "phd", "b.s.", "m.s.", "b.tech", "m.tech", "b.a.", "m.a.", "bba", "mba", "associate degree", "diploma"]
    text_lower = text.lower()
    found = []
    
    lines = text.split("\n")
    in_edu_section = False
    
    for line in lines:
        if "education" in line.lower() or "academic" in line.lower():
            in_edu_section = True
            continue
        if in_edu_section and any(sec in line.lower() for sec in ["experience", "projects", "skills", "certifications"]):
            in_edu_section = False
            
        if in_edu_section or any(d in line.lower() for d in degrees):
            cleaned_line = line.strip()
            if cleaned_line and len(cleaned_line) > 5 and len(cleaned_line) < 150:
                found.append(cleaned_line)
                if len(found) >= 3:
                    break
                    
    return "; ".join(found) if found else "Degree / Relevant Coursework listed"

def parse_experience(text):
    """
    Extracts experience section highlights or years of experience.
    """
    exp_pattern = r'(\d+\+?\s*(?:years?|yrs?)\s*(?:of)?\s*(?:experience|exp)?)'
    matches = re.findall(exp_pattern, text, re.IGNORECASE)
    
    summary = []
    if matches:
        summary.append(matches[0])
        
    lines = text.split("\n")
    in_exp = False
    for line in lines:
        if "experience" in line.lower() or "work history" in line.lower() or "employment" in line.lower():
            in_exp = True
            continue
        if in_exp and any(sec in line.lower() for sec in ["education", "projects", "skills", "certifications"]):
            in_exp = False
        if in_exp:
            line_str = line.strip()
            if line_str and len(line_str) > 10 and not line_str.startswith("-"):
                summary.append(line_str[:100])
                if len(summary) >= 4:
                    break
                    
    return " | ".join(summary) if summary else "Professional experience detailed in resume."

def parse_projects_certifications(text):
    projects = []
    certs = []
    lines = text.split("\n")
    
    for line in lines:
        l = line.lower()
        if "certified" in l or "certification" in l or "certificate" in l or "aws certified" in l:
            certs.append(line.strip()[:80])
        elif "project:" in l or "developed" in l or "built" in l:
            projects.append(line.strip()[:80])
            
    return (
        "; ".join(projects[:3]) if projects else "Projects documented",
        "; ".join(certs[:3]) if certs else "Certifications documented"
    )

def parse_resume(filepath):
    """
    Main entry point: extracts text from file and parses structured candidate details.
    """
    raw_text = extract_text_from_file(filepath)
    if not raw_text.strip():
        # Fallback if text extraction fails
        raw_text = f"Sample text for file {os.path.basename(filepath)}"
        
    name = parse_name(raw_text)
    email = parse_email(raw_text)
    phone = parse_phone(raw_text)
    skills = parse_skills(raw_text)
    education = parse_education(raw_text)
    experience = parse_experience(raw_text)
    projects, certs = parse_projects_certifications(raw_text)
    
    return {
        "raw_text": raw_text,
        "name": name,
        "email": email,
        "phone": phone,
        "skills": skills,
        "education": education,
        "experience": experience,
        "projects": projects,
        "certifications": certs
    }
